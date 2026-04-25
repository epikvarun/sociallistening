import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "..", "..", "data", "social_listening_posts.xlsx")

HEADERS = ["Date Found", "Platform", "Topic", "Post Text", "URL", "Epik's Take"]
COLUMN_WIDTHS = [20, 12, 18, 60, 60, 40]

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _ensure_workbook() -> openpyxl.Workbook:
    os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)
    if os.path.exists(EXCEL_PATH):
        return openpyxl.load_workbook(EXCEL_PATH)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Posts"
    for col_idx, (header, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"
    wb.save(EXCEL_PATH)
    return wb


def _parse_slack_message(message: str) -> dict:
    """Extract structured fields from the standard Slack message format."""
    platform_match = re.search(r"\[(\w+)\]", message)
    topic_match = re.search(r"topic:\s*`([^`]+)`", message)
    text_match = re.search(r"^>\s*(.+)", message, re.MULTILINE)
    url_match = re.search(r"(https?://\S+)", message)
    take_match = re.search(r'Epik\'s take:\s*"([^"]+)"', message)

    return {
        "platform": platform_match.group(1).capitalize() if platform_match else "",
        "topic": topic_match.group(1) if topic_match else "",
        "text": text_match.group(1).strip() if text_match else "",
        "url": url_match.group(1) if url_match else "",
        "take": take_match.group(1) if take_match else "",
    }


def append_post(message: str) -> None:
    """Parse a Slack-formatted message and append one row to the Excel file."""
    fields = _parse_slack_message(message)
    if not fields["url"]:
        return

    wb = _ensure_workbook()
    ws = wb["Posts"]

    next_row = ws.max_row + 1
    values = [
        datetime.now().strftime("%Y-%m-%d %H:%M"),
        fields["platform"],
        fields["topic"],
        fields["text"],
        fields["url"],
        fields["take"],
    ]

    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=next_row, column=col_idx, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if col_idx == 5:  # URL column — make it a hyperlink
            cell.hyperlink = value
            cell.font = Font(color="0563C1", underline="single")

    ws.row_dimensions[next_row].height = 40
    wb.save(EXCEL_PATH)
